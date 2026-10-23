#!/usr/bin/env python3
# Benjamini-Hochberg FDR over the step 2-6 tests, within each family and experiment-wide (85 tests).
# Usage: python3 scripts/07_FDR_correction.py config.sh
import os
import sys
import numpy as np


def read_config(path):
    cfg = {}
    with open(path) as fh:
        for line in fh:
            line = line.split("#", 1)[0].strip()
            if "=" in line:
                k, v = line.split("=", 1)
                cfg[k.strip()] = os.path.expanduser(v.strip().strip("\"'"))
    return cfg


cfg_path = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("PIPELINE_CONFIG", "")
if not os.path.isfile(cfg_path):
    sys.exit("pass the config file as the first argument or set PIPELINE_CONFIG")
OUT = read_config(cfg_path)["OUT_DIR"]
os.makedirs(OUT, exist_ok=True)
M3  = os.path.join(OUT, "PRS_pBAG_linear_results.txt")

def fdr_bh(p):
    p = np.asarray(p, dtype=float); n = len(p)
    order = np.argsort(p); ranks = np.empty_like(order)
    ranks[order] = np.arange(1, n + 1)
    adj = p * n / ranks
    adj_sorted = adj[order]
    for i in range(n - 2, -1, -1):
        adj_sorted[i + 1] = min(adj_sorted[i + 1], 1.0)
        adj_sorted[i]     = min(adj_sorted[i], adj_sorted[i + 1])
    adj_sorted[-1] = min(adj_sorted[-1], 1.0)
    out = np.empty(n); out[order] = adj_sorted
    return out


rows = []   # (method, label, raw_p)

# Apart from M3, the p-values below are copied from the step 2, 5 and 6 outputs of the
# reference run, so they must be updated by hand after a re-run.

# M1: PRSice best-PT p-values (8 GWAS x 4 contrasts)
M1 = {
    "Epilepsy":        {"Syndrome": 1.77655e-05, "DRE": 0.431862,
                        "Convulsions": 0.136059,  "VPA": 0.290732},
    "Intelligence":    {"Syndrome": 0.428233,    "DRE": 0.00170035,
                        "Convulsions": 0.00129029,"VPA": 0.696364},
    "Longevity":       {"Syndrome": 0.186247,    "DRE": 0.0904619,
                        "Convulsions": 0.180848,  "VPA": 0.00300832},
    "AD":              {"Syndrome": 0.582445,    "DRE": 0.116283,
                        "Convulsions": 0.207875,  "VPA": 0.0114548},
    "Frailty":         {"Syndrome": 0.206639,    "DRE": 0.212172,
                        "Convulsions": 0.425797,  "VPA": 0.120776},
    "Brain Age (GM)":  {"Syndrome": 0.0174842,   "DRE": 0.492195,
                        "Convulsions": 0.147377,  "VPA": 0.0396162},
    "Brain Age (GWM)": {"Syndrome": 0.0312526,   "DRE": 0.45964,
                        "Convulsions": 0.0890036, "VPA": 0.216405},
    "Brain Age (WM)":  {"Syndrome": 0.0518639,   "DRE": 0.234693,
                        "Convulsions": 0.153357,  "VPA": 0.635745},
}
for g, fs in M1.items():
    for f, p in fs.items():
        rows.append(("M1: PRS group comparison", f"{g} x {f}", p))

# M3: best PRS->pBAG p per GWAS (read from M3 results)
m3_best = {}
try:
    with open(M3) as fh:
        next(fh)
        for line in fh:
            parts = line.strip().split("\t")
            g, p = parts[0], float(parts[-1])
            if g not in m3_best or p < m3_best[g]:
                m3_best[g] = p
except Exception:
    m3_best = {"Brain Age (GWM)": 1.521e-3, "Frailty": 0.0907,
               "Epilepsy": 0.156, "Longevity": 0.203, "Intelligence": 0.359}
for g, p in m3_best.items():
    rows.append(("M3: Individual PRS->pBAG regression", g, p))

# M4: hierarchical block F-test + PRS coefficients in full model
for label, p in [("Block 1: Demographics", 5.43e-01),
                 ("Block 2: Epilepsy factors", 6.90e-02),
                 ("Block 3: PRS (genetics)", 2.98e-04)]:
    rows.append(("M4: Hierarchical block F-test", label, p))

M4 = {"PRS Epilepsy": 0.170754, "PRS Intelligence": 0.720110,
      "PRS Longevity": 0.080755, "PRS Frailty": 0.001279,
      "PRS Brain Age (GWM)": 0.013345}
for label, p in M4.items():
    rows.append(("M4: Hierarchical PRS coefficient", label, p))

# M5a: stratification ANOVA + linear trend
M5a = [
    ("Epilepsy",        "Tertile",  0.0418, 0.0139),
    ("Epilepsy",        "Quartile", 0.345,  0.0801),
    ("Intelligence",    "Tertile",  0.924,  0.991),
    ("Intelligence",    "Quartile", 0.216,  0.696),
    ("Longevity",       "Tertile",  0.545,  0.291),
    ("Longevity",       "Quartile", 0.812,  0.388),
    ("Frailty",         "Tertile",  0.214,  0.0821),
    ("Frailty",         "Quartile", 0.254,  0.168),
    ("Brain Age (GWM)", "Tertile",  0.03,   0.0152),
    ("Brain Age (GWM)", "Quartile", 0.0109, 0.00259),
]
for g, bin_, ap, tp in M5a:
    rows.append(("M5a: Stratification ANOVA", f"{g} {bin_}", ap))
    rows.append(("M5a: Stratification trend",  f"{g} {bin_}", tp))

# M5b: PRS x clinical interactions
M5b = [
    ("Epilepsy",        "Syndrome",    0.835),
    ("Epilepsy",        "DRE",         0.577),
    ("Epilepsy",        "Convulsions", 0.149),
    ("Epilepsy",        "VPA",         0.824),
    ("Intelligence",    "Syndrome",    0.942),
    ("Intelligence",    "DRE",         0.941),
    ("Intelligence",    "Convulsions", 0.762),
    ("Intelligence",    "VPA",         0.152),
    ("Longevity",       "Syndrome",    0.653),
    ("Longevity",       "DRE",         0.866),
    ("Longevity",       "Convulsions", 0.909),
    ("Longevity",       "VPA",         0.36),
    ("Frailty",         "Syndrome",    0.767),
    ("Frailty",         "DRE",         0.934),
    ("Frailty",         "Convulsions", 0.259),
    ("Frailty",         "VPA",         0.516),
    ("Brain Age (GWM)", "Syndrome",    0.781),
    ("Brain Age (GWM)", "DRE",         0.475),
    ("Brain Age (GWM)", "Convulsions", 0.505),
    ("Brain Age (GWM)", "VPA",         0.0587),
]
for g, f, p in M5b:
    rows.append(("M5b: PRSxclinical interaction", f"{g} x {f}", p))


methods = [r[0] for r in rows]
labels  = [r[1] for r in rows]
raw_p   = np.array([r[2] for r in rows])

# Within-family FDR
fam_q = np.full(len(raw_p), np.nan)
for m in dict.fromkeys(methods):
    idx = [i for i, mm in enumerate(methods) if mm == m]
    fam_q[idx] = fdr_bh(raw_p[idx])

# Experiment-wide FDR
exp_q = fdr_bh(raw_p)


def fmt(p): return f"{p:.2e}" if p < 0.001 else f"{p:.4f}"
def star(p): return "***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else ""

out_txt = os.path.join(OUT, "FDR_correction_summary.txt")
with open(out_txt, "w") as fh:
    fh.write("EXPERIMENT-WIDE FDR (Benjamini-Hochberg)\n")
    fh.write(f"Total tests: {len(raw_p)}\n\n")
    fh.write("Method summary:\n")
    for m in dict.fromkeys(methods):
        idx = [i for i, mm in enumerate(methods) if mm == m]
        fh.write(f"  {m:45s} tests={len(idx):3d}  "
                 f"raw<0.05={sum(raw_p[i]<0.05 for i in idx)}  "
                 f"famFDR<0.05={sum(fam_q[i]<0.05 for i in idx)}  "
                 f"expFDR<0.05={sum(exp_q[i]<0.05 for i in idx)}\n")
    fh.write("\n")
    fh.write(f"{'Method':<45s}  {'Test':<30s}  {'Raw p':>10s}  {'Fam-FDR':>10s}  {'Exp-FDR':>10s}  Sig\n")
    fh.write("-" * 120 + "\n")
    cur = None
    for i in range(len(rows)):
        if methods[i] != cur:
            if cur is not None: fh.write("\n")
            cur = methods[i]
        fh.write(f"{methods[i]:<45s}  {labels[i]:<30s}  "
                 f"{fmt(raw_p[i]):>10s}  {fmt(fam_q[i]):>10s}  "
                 f"{fmt(exp_q[i]):>10s}  {star(exp_q[i])}\n")
    surv = [(methods[i], labels[i], raw_p[i], exp_q[i])
            for i in range(len(rows)) if exp_q[i] < 0.05]
    fh.write("\nSurvivors (exp-FDR q < 0.05):\n")
    for m, l, rp, q in sorted(surv, key=lambda x: x[3]):
        fh.write(f"  {m:45s}  {l:<30s}  raw={fmt(rp)}  FDR={fmt(q)}\n")
print(f"Saved: {out_txt}")

# Excel (if openpyxl is around)
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "FDR correction"
    for j, h in enumerate(["Method","Test","Raw p","Family-FDR q",
                           "Experiment-FDR q","Sig"], 1):
        c = ws.cell(row=1, column=j, value=h); c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for i, _ in enumerate(rows):
        ws.cell(row=i+2, column=1, value=methods[i])
        ws.cell(row=i+2, column=2, value=labels[i])
        ws.cell(row=i+2, column=3, value=round(raw_p[i], 6))
        ws.cell(row=i+2, column=4, value=round(fam_q[i], 6))
        ws.cell(row=i+2, column=5, value=round(exp_q[i], 6))
        ws.cell(row=i+2, column=6, value=star(exp_q[i]))
        if exp_q[i] < 0.05:
            for j in range(1, 7):
                ws.cell(row=i+2, column=j).fill = PatternFill("solid", fgColor="C6EFCE")
    for col, w in zip("ABCDEF", [48, 35, 14, 16, 20, 10]):
        ws.column_dimensions[col].width = w
    xlsx_path = os.path.join(OUT, "FDR_correction_summary.xlsx")
    wb.save(xlsx_path); print(f"Saved: {xlsx_path}")
except ImportError:
    print("(openpyxl not installed, skipped Excel)")
